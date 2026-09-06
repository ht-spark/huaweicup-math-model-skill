#!/usr/bin/env python3
"""Render a Huawei Cup Markdown report as a Word-native DOCX.

Pandoc performs Markdown and TeX-math conversion.  This script then applies
competition-paper formatting that Pandoc does not express by itself:

* display equations remain editable Office Math (OMML), with source ``\tag``
  labels placed at the right margin;
* every semantic table uses a three-line layout (top rule, header rule,
  bottom rule) with no vertical or internal grid lines.
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Optional
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm


DISPLAY_MATH_RE = re.compile(r"(?<!\\)\$\$(.*?)(?<!\\)\$\$", re.DOTALL)
TAG_RE = re.compile(r"\\tag\s*\{([^{}]+)\}")
FULLWIDTH_TRAILING_PUNCTUATION_RE = re.compile(r"[，；。]\s*$")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert a Markdown report to DOCX with OMML equations and three-line tables."
    )
    parser.add_argument("source", type=Path, help="Source Markdown report")
    parser.add_argument("-o", "--output", type=Path, required=True, help="Output DOCX path")
    parser.add_argument(
        "--reference-doc",
        type=Path,
        help="Optional official/template DOCX passed to pandoc as --reference-doc",
    )
    parser.add_argument(
        "--pandoc",
        type=Path,
        help="Pandoc executable; defaults to the pandoc found on PATH",
    )
    return parser.parse_args()


def normalize_display_math(markdown: str) -> tuple[str, list[Optional[str]]]:
    """Remove TeX tags/punctuation that Pandoc cannot typeset as Word numbering."""

    labels: list[Optional[str]] = []

    def replace(match: re.Match[str]) -> str:
        body = match.group(1).strip()
        found = TAG_RE.findall(body)
        if len(found) > 1:
            raise ValueError("One display-math block contains more than one \\tag label.")
        label = found[0].strip() if found else None
        body = TAG_RE.sub("", body).rstrip()
        body = body.translate(str.maketrans("", "", "，；。"))
        body = FULLWIDTH_TRAILING_PUNCTUATION_RE.sub("", body).rstrip()
        labels.append(label)
        return f"$$\n{body}\n$$"

    return DISPLAY_MATH_RE.sub(replace, markdown), labels


def _replace_child(parent, tag: str, child) -> None:
    existing = parent.find(qn(tag))
    if existing is not None:
        parent.remove(existing)
    parent.append(child)


def _border(edge: str, value: str, size: int = 0) -> OxmlElement:
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), value)
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), "0")
    element.set(qn("w:color"), "000000" if value == "single" else "auto")
    return element


def apply_three_line_table(table) -> None:
    """Apply a Chinese academic three-line table using direct OOXML borders."""

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    try:
        table.style = "Table Normal"
    except KeyError:
        pass

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    borders.append(_border("top", "single", 12))       # 1.5 pt
    borders.append(_border("left", "nil"))
    borders.append(_border("bottom", "single", 12))    # 1.5 pt
    borders.append(_border("right", "nil"))
    borders.append(_border("insideH", "nil"))
    borders.append(_border("insideV", "nil"))
    _replace_child(tbl_pr, "w:tblBorders", borders)

    look = OxmlElement("w:tblLook")
    look.set(qn("w:val"), "0000")
    look.set(qn("w:firstRow"), "0")
    look.set(qn("w:lastRow"), "0")
    look.set(qn("w:firstColumn"), "0")
    look.set(qn("w:lastColumn"), "0")
    look.set(qn("w:noHBand"), "1")
    look.set(qn("w:noVBand"), "1")
    _replace_child(tbl_pr, "w:tblLook", look)

    if not table.rows:
        return

    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            old_borders = tc_pr.find(qn("w:tcBorders"))
            if old_borders is not None:
                tc_pr.remove(old_borders)
            shading = tc_pr.find(qn("w:shd"))
            if shading is not None:
                tc_pr.remove(shading)

            if row_index == 0:
                cell_borders = OxmlElement("w:tcBorders")
                cell_borders.append(_border("bottom", "single", 6))  # 0.75 pt
                tc_pr.append(cell_borders)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def apply_default_page_layout(document: Document) -> None:
    """Use deterministic A4 geometry when no official reference DOCX is supplied."""

    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(25.4)
        section.bottom_margin = Mm(25.4)
        section.left_margin = Mm(25.4)
        section.right_margin = Mm(25.4)


def _tab_run() -> OxmlElement:
    run = OxmlElement("w:r")
    run.append(OxmlElement("w:tab"))
    return run


def _number_run(label: str) -> OxmlElement:
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    clean_label = label.strip()
    if clean_label.startswith("(") and clean_label.endswith(")"):
        clean_label = clean_label[1:-1].strip()
    text.text = f"({clean_label})"
    run.append(text)
    return run


def apply_equation_numbers(document: Document, labels: list[Optional[str]]) -> int:
    """Place display-math labels at the right margin while retaining OMML."""

    math_paragraphs = []
    for paragraph in document.element.body.iter(qn("w:p")):
        math_para = paragraph.find(qn("m:oMathPara"))
        if math_para is not None:
            math_paragraphs.append((paragraph, math_para))

    if len(math_paragraphs) != len(labels):
        raise RuntimeError(
            "Display-equation count changed during conversion: "
            f"Markdown={len(labels)}, DOCX={len(math_paragraphs)}."
        )

    if not document.sections:
        raise RuntimeError("The generated DOCX has no page section information.")
    section = document.sections[0]
    if (
        section.page_width is not None
        and section.left_margin is not None
        and section.right_margin is not None
    ):
        content_width_twips = int(
            (section.page_width - section.left_margin - section.right_margin) / 635
        )
    else:
        # Pandoc's default reference document may inherit section dimensions
        # instead of writing them into sectPr.  9026 twips is the writable
        # width of A4 paper with Word's 1-inch left and right margins.
        content_width_twips = 9026
    center_position = content_width_twips // 2
    right_position = content_width_twips

    numbered = 0
    for (paragraph, math_para), label in zip(math_paragraphs, labels):
        if label is None:
            continue

        math_nodes = [child for child in math_para if child.tag == qn("m:oMath")]
        if not math_nodes:
            raise RuntimeError("A display-equation paragraph contains no Office Math object.")

        paragraph_properties = paragraph.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        center_tab = OxmlElement("w:tab")
        center_tab.set(qn("w:val"), "center")
        center_tab.set(qn("w:pos"), str(center_position))
        tabs.append(center_tab)
        right_tab = OxmlElement("w:tab")
        right_tab.set(qn("w:val"), "right")
        right_tab.set(qn("w:pos"), str(right_position))
        tabs.append(right_tab)
        _replace_child(paragraph_properties, "w:tabs", tabs)

        math_index = paragraph.index(math_para)
        paragraph.remove(math_para)
        paragraph.insert(math_index, _tab_run())
        insert_index = math_index + 1
        for math_node in math_nodes:
            math_para.remove(math_node)
            paragraph.insert(insert_index, math_node)
            insert_index += 1
        paragraph.insert(insert_index, _tab_run())
        paragraph.insert(insert_index + 1, _number_run(label))
        numbered += 1

    return numbered


def validate_docx(
    path: Path,
    expected_tables: int,
    expected_numbered: int,
    expect_office_math: bool,
) -> tuple[int, int]:
    """Check that formulas are OMML and table border contracts are present."""

    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml")

    if b"\\tag{" in xml or b"$$" in xml:
        raise RuntimeError("Raw TeX equation markers remain in the generated DOCX.")

    document = Document(path)
    if len(document.tables) != expected_tables:
        raise RuntimeError(
            f"Table count changed after formatting: expected={expected_tables}, "
            f"actual={len(document.tables)}."
        )

    for index, table in enumerate(document.tables, start=1):
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        if borders is None:
            raise RuntimeError(f"Table {index} has no direct border definition.")
        required = {
            "top": ("single", "12"),
            "bottom": ("single", "12"),
            "left": ("nil", "0"),
            "right": ("nil", "0"),
            "insideH": ("nil", "0"),
            "insideV": ("nil", "0"),
        }
        for edge, (value, size) in required.items():
            item = borders.find(qn(f"w:{edge}"))
            if item is None or item.get(qn("w:val")) != value:
                raise RuntimeError(f"Table {index} violates the three-line rule at {edge}.")
            if item.get(qn("w:sz")) != size:
                raise RuntimeError(f"Table {index} has an unexpected {edge} line width.")

        if not table.rows:
            continue
        for cell in table.rows[0].cells:
            cell_borders = cell._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
            header_bottom = (
                cell_borders.find(qn("w:bottom")) if cell_borders is not None else None
            )
            if (
                header_bottom is None
                or header_bottom.get(qn("w:val")) != "single"
                or header_bottom.get(qn("w:sz")) != "6"
            ):
                raise RuntimeError(f"Table {index} has no 0.75 pt header rule.")

    office_math_count = len(document.element.xpath(".//m:oMath"))
    equation_number_count = 0
    for paragraph in document.element.body.iter(qn("w:p")):
        if paragraph.find(qn("m:oMath")) is not None:
            texts = paragraph.findall(qn("w:r"))
            if any(
                "(" in "".join(run.itertext()) and ")" in "".join(run.itertext())
                for run in texts
            ):
                equation_number_count += 1

    if expect_office_math and office_math_count == 0:
        raise RuntimeError("The generated DOCX contains no editable Office Math equations.")
    if equation_number_count != expected_numbered:
        raise RuntimeError(
            "Equation-number count mismatch: "
            f"expected={expected_numbered}, actual={equation_number_count}."
        )

    return office_math_count, equation_number_count


def main() -> int:
    args = parse_args()
    source = args.source.resolve()
    output = args.output.resolve()

    if not source.is_file():
        raise FileNotFoundError(f"Markdown source not found: {source}")
    if source.suffix.lower() not in {".md", ".markdown"}:
        raise ValueError("The source file must be Markdown (.md or .markdown).")
    if output.suffix.lower() != ".docx":
        raise ValueError("The output file must use the .docx extension.")

    reference_doc = args.reference_doc.resolve() if args.reference_doc else None
    if reference_doc is not None and not reference_doc.is_file():
        raise FileNotFoundError(f"Reference DOCX not found: {reference_doc}")

    pandoc = str(args.pandoc) if args.pandoc else shutil.which("pandoc")
    if not pandoc:
        raise RuntimeError("Pandoc was not found. Install Pandoc or pass --pandoc PATH.")

    markdown = source.read_text(encoding="utf-8")
    normalized_markdown, labels = normalize_display_math(markdown)
    expected_numbered = sum(label is not None for label in labels)

    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="huaweicup-docx-") as temp_dir:
        normalized_source = Path(temp_dir) / source.name
        normalized_source.write_text(normalized_markdown, encoding="utf-8")
        command = [
            pandoc,
            str(normalized_source),
            "--from=markdown+tex_math_dollars",
            "--to=docx",
            "--standalone",
            f"--resource-path={source.parent}",
            f"--output={output}",
        ]
        if reference_doc is not None:
            command.append(f"--reference-doc={reference_doc}")
        subprocess.run(command, cwd=source.parent, check=True)

    document = Document(output)
    if reference_doc is None:
        apply_default_page_layout(document)
    table_count = len(document.tables)
    for table in document.tables:
        apply_three_line_table(table)
    numbered_count = apply_equation_numbers(document, labels)
    document.save(output)

    office_math_count, verified_numbered = validate_docx(
        output,
        expected_tables=table_count,
        expected_numbered=expected_numbered,
        expect_office_math=bool(labels) or bool(re.search(r"(?<!\\)\$(?!\$).+?(?<!\\)\$", markdown)),
    )
    print(f"DOCX: {output}")
    print(f"Three-line tables: {table_count}")
    print(f"Editable Office Math objects: {office_math_count}")
    print(f"Right-aligned equation numbers: {verified_numbered}")
    if numbered_count != verified_numbered:
        raise RuntimeError("Equation numbering changed during the final save.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (OSError, RuntimeError, ValueError, subprocess.CalledProcessError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
